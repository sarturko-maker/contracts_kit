"""Invented export integration tests; no external service or model calls."""

import csv
import json
import subprocess
import sys
from pathlib import Path

from test_place import KitFixture, rows, write_rows
from review_table import COLUMNS


def answers_from_card(card):
    groups = {"document": ("q1",), "parties": ("q2",), "execution": ("q3",),
              "term": ("q4",), "trade_scope": ("q7",), "group_scope": ("q8",),
              "links": ("q6",), "precedence": ("q5_precedence",), "parts": ("q5",),
              "gaps": ("q9", "q10")}
    result = {col: "\n".join(f"{k}: {v}" for k, v in card.items() if any(k.startswith(p) for p in prefixes))
              for col, prefixes in groups.items()}
    return result


class ReviewTableTest(KitFixture):
    def setUp(self):
        super().setUp()
        self.cards = {p.stem: json.loads(p.read_text()) for p in (self.root / "work/cards").glob("*.json")}
        self.inventory = rows(self.root / "work/inventory.csv")
        self.table_rows = [{"file_name": r["file_name"], **answers_from_card(self.cards[r["doc_id"]])}
                           for r in self.inventory if r["doc_id"].isdigit() and r["readable"] == "yes"]
        self.table = self.pile / "Review_Table.csv"
        self.save_table()

    def save_table(self, delimiter=","):
        with self.table.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(self.table_rows[0]), delimiter=delimiter)
            writer.writeheader()
            writer.writerows(self.table_rows)

    def run_review(self, *args, success=True):
        result = subprocess.run([sys.executable, str(self.root / "scripts/review_table.py"), *args],
                                cwd=self.root, text=True, capture_output=True)
        self.assertEqual(result.returncode == 0, success, result.stdout + result.stderr)
        return result.stdout

    def import_table(self):
        return self.run_review("--import", str(self.table))

    def packet(self, doc="001"):
        active = json.loads((self.root / "work/review-table/active.json").read_text())
        return json.loads((self.root / "work/review-table/rows" / active["rows"][doc] / f"{doc}.json").read_text())

    def assignments(self, holding=False):
        decisions = [{"doc_id": doc, "account": self.erp_accounts[int(doc) % 2],
                      "companies_found": "Invented party", "basis": "parties cell", "confidence": "sure",
                      "note": "invented assignment"} for doc in sorted(self.cards)]
        if holding:
            decisions[0]["account"] = "_not-sure/Unknown party"
        path = self.root / "work/review-table/assignments.json"
        path.write_text(json.dumps(decisions))
        self.run_review("--assign", str(path))
        return decisions

    def test_prepare_registers_control_file_without_numbering_it(self):
        self.command("prepare.py", str(self.pile), "--account-column", "customer_account", "--side", "customers")
        self.assertEqual(self.inventory, rows(self.root / "work/inventory.csv"))
        pointer = json.loads((self.root / "work/review-source.json").read_text())
        self.assertEqual(str(self.table), pointer["source_path"])
        self.assertIn("9 rows imported", self.run_review("--import"))

    def test_raw_index_accepts_nested_quotes_blanks_and_no_attestation(self):
        self.table_rows[0]["term"] = '  The term was extended: "defined as "Initial Term" here".  '
        self.table_rows[0]["execution"] = "Reported signed; no image attestation supplied."
        self.table_rows[0]["gaps"] = ""
        self.save_table()
        self.import_table()
        self.assertEqual(self.table_rows[0], self.packet()["answers"])
        self.assertEqual(["gaps: blank"], self.packet()["flags"])
        self.assertIn("9 index rows ready", self.run_review("--status"))
        self.assertFalse((self.root / "work/cards/001.json").exists())
        self.assertIn("0 changed", self.import_table())
        self.assertIn("Initial Term", self.run_review("--index"))
        self.assertFalse((self.root / "work/review-table/source-checks.jsonl").exists())

    def test_direct_index_and_judgments_render_without_cards_or_source_reads(self):
        self.table_rows[0]["gaps"] = ""
        self.save_table()
        self.import_table()
        self.assignments(holding=True)
        for account in self.erp_accounts:
            self.judge(account, self.sorted_docs(account))  # invented decisions test rendering, not accuracy
        self.command("place.py", "--all", "--visuals")
        corpus = rows(self.root / "out/customers/CORPUS.csv")
        self.assertEqual({r["doc_id"] for r in self.inventory if r["doc_id"].isdigit()},
                         {r["doc_id"] for r in corpus})
        doc = next(r for r in corpus if r["doc_id"] == "001")
        self.assertEqual("", doc["review_gaps"])
        self.assertEqual(self.table_rows[0]["execution"], doc["review_execution"])
        self.assertEqual("none", doc["source_checks"])
        self.assertEqual("unassessed (no account)", doc["status"])
        folder = self.root / "out/customers/_not-sure/Unknown party"
        self.assertTrue((folder / doc["filed_as"]).exists())
        self.assertTrue((folder / "documents.csv").exists())
        self.assertTrue((folder / "README.md").exists())
        for account in self.erp_accounts:
            for name in ("README.md", "ANALYSIS.md", "position.html"):
                self.assertIn("Evidence basis: Review_Table extraction",
                              (self.root / "out/customers" / account / name).read_text())
            self.assertNotIn("no start date found", (self.root / "out/customers" / account / "ANALYSIS.md").read_text())
        self.assertFalse(list((self.root / "work/cards").glob("*.json")))
        self.assertFalse((self.root / "work/review-table/source-checks.jsonl").exists())
        self.command("visualise.py", "--all", "--analysis")
        self.assertTrue((self.root / "out/INDEX.html").exists())
        self.assertIn("Use /analyse", self.command("sort.py", success=False))

    def test_semicolon_csv_keeps_multiline_answers(self):
        self.save_table(delimiter=";")
        self.assertIn("9 rows imported", self.import_table())

    def test_xlsx_values_and_multiple_sheet_selection(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "Answers"
        sheet.append(list(self.table_rows[0]))
        for row in self.table_rows:
            sheet.append(list(row.values()))
        path = self.pile / "export.xlsx"
        book.save(path)
        self.assertIn("9 rows imported", self.run_review("--import", str(path)))
        book.create_sheet("Notes").append(["export notes"])
        book.save(path)
        self.assertIn("explicit --sheet", self.run_review("--import", str(path), success=False))
        self.run_review("--import", str(path), "--sheet", "Answers")
        self.run_review("--import")  # recorded sheet survives the next stage run
        sheet.cell(2, 2).value = '=HYPERLINK("https://example.invalid", "answer")'
        book.save(path)
        self.assertIn("formulas", self.run_review("--import", str(path), "--sheet", "Answers", success=False))

    def test_missing_column_or_row_does_not_replace_prior_cards(self):
        original = (self.root / "work/cards/001.json").read_bytes()
        self.table_rows.pop()
        self.save_table()
        self.assertIn("no row", self.run_review("--import", str(self.table), success=False))
        self.assertEqual(original, (self.root / "work/cards/001.json").read_bytes())
        self.table.write_text("file_name,document\nexample.pdf,Master\n")
        self.assertIn("missing columns", self.run_review("--import", str(self.table), success=False))

    def test_duplicate_or_contradictory_identity_is_rejected(self):
        self.table_rows.append(dict(self.table_rows[0]))
        self.save_table()
        self.assertIn("Multiple Review_Table rows", self.run_review("--import", str(self.table), success=False))
        self.table_rows.pop()
        for row in self.table_rows:
            row["sha256"] = "wrong"
        self.save_table()
        self.assertIn("contradictory", self.run_review("--import", str(self.table), success=False))

    def test_ambiguous_filename_requires_disambiguating_metadata(self):
        inventory = list(self.inventory)
        first = next(r for r in inventory if r["doc_id"] == "001")
        second = next(r for r in inventory if r["doc_id"] == "002")
        second["file_name"] = first["file_name"]
        write_rows(self.root / "work/inventory.csv", inventory, list(inventory[0]))
        self.assertIn("ambiguous", self.run_review("--import", str(self.table), success=False))

    def test_source_check_requires_reason_and_preserves_actual_scope_and_finding(self):
        self.import_table()
        request = self.root / "work/review-table/request.json"
        request.write_text(json.dumps({"doc_id":"001", "mode":"text", "location":"page 2"}))
        self.assertIn("requires reason", self.run_review("--request-check", str(request), success=False))
        request.write_text(json.dumps({"doc_id":"001", "mode":"text", "location":"page 2",
                                      "reason":"Resolve material priority conflict"}))
        self.run_review("--request-check", str(request))
        log = self.root / "work/review-table/source-checks.jsonl"
        event = json.loads(log.read_text().splitlines()[0])
        self.assertIn("1 pending", self.run_review("--status"))
        result = self.root / "work/review-table/result.json"
        result.write_text(json.dumps({"check_id":event["check_id"], "mode":"text",
                                    "location":"page 2, clause 11.1", "finding":"Invented test outcome"}))
        self.run_review("--finish-check", str(result))
        self.assertIn("1 completed, 0 pending", self.run_review("--status"))
        self.assertIn("Invented test outcome", self.run_review("--index"))
        self.assertIn("already", self.run_review("--finish-check", str(result), success=False))
        self.table_rows[0]["gaps"] += "changed"
        self.save_table()
        self.import_table()
        self.assertIn("0 requested", self.run_review("--status"))
        self.assertEqual(2, len(log.read_text().splitlines()))  # historical checks preserved, not reused

    def test_changed_source_or_export_blocks_reuse(self):
        self.import_table()
        source = Path(next(r["original_path"] for r in self.inventory if r["doc_id"] == "001"))
        original = source.read_bytes()
        source.write_bytes(original + b"changed")
        self.assertIn("source changed", self.run_review("--status", success=False))
        source.write_bytes(original)
        self.table.write_text(self.table.read_text(encoding="utf-8-sig") + "\n")
        self.assertIn("changed or disappeared", self.run_review("--status", success=False))

    def test_invalid_assignment_cannot_overwrite_matching(self):
        self.import_table()
        decisions = self.assignments()
        log = self.root / "work/logs/sort.csv"
        before = log.read_bytes()
        path = self.root / "work/review-table/assignments.json"
        decisions[0]["account"] = "Not an ERP account"
        path.write_text(json.dumps(decisions))
        self.assertIn("not an ERP row", self.run_review("--assign", str(path), success=False))
        self.assertEqual(before, log.read_bytes())
        path.write_text(json.dumps(decisions[1:]))
        self.assertIn("every selected", self.run_review("--assign", str(path), success=False))
        self.assertEqual(before, log.read_bytes())

    def test_scoped_reassignment_preserves_other_rows_and_invalidates_shared_judgment(self):
        self.import_table()
        decisions = self.assignments()
        for account in self.erp_accounts:
            self.judge(account, self.sorted_docs(account))
        before = rows(self.root / "work/logs/sort.csv")
        account = self.erp_accounts[0]
        selected = [r for r in decisions if r["account"] == account]
        selected[0]["account"] = self.erp_accounts[1]
        path = self.root / "work/review-table/scoped.json"
        path.write_text(json.dumps(selected))
        self.run_review("--assign", str(path), "--account", account)
        after = rows(self.root / "work/logs/sort.csv")
        self.assertEqual(len(before), len(after))
        for row in before:
            if row["doc_id"] not in {r["doc_id"] for r in selected}:
                self.assertIn(row, after)
        self.assertFalse(list((self.root / "work/placements").glob("*.csv")))
        self.assertTrue(list((self.root / "work/history").glob("review-match-*/*.csv")))

    def test_changed_row_archives_dependent_reports(self):
        self.import_table()
        report = self.root / "out/INDEX.md"
        report.parent.mkdir(parents=True, exist_ok=True)
        report.write_text("prior report")
        placement = self.root / "work/placements/Example.md"
        placement.parent.mkdir(parents=True, exist_ok=True)
        placement.write_text("prior position")
        self.table_rows[0]["gaps"] += " changed"
        self.save_table()
        self.import_table()
        self.assertFalse(report.exists())
        self.assertFalse(placement.exists())
        self.assertTrue(list((self.root / "work/history").glob("review-*/out/INDEX.md")))
