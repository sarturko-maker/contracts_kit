"""Import Review_Table as the analysis index. No model conversion or source reading.

Cells stay verbatim. The orchestrator assigns accounts and judges directly from the index;
only material doubts warrant a separately logged, targeted source check.
"""

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from kit_common import (CARD_KEYS, INVENTORY_CSV, WORK, OUT, WORK_FILES,
                        SORT_LOG_COLUMNS, read_csv, read_json, write_csv, write_json,
                        safe_folder_name, fail, say)

COLUMNS = ("document", "parties", "execution", "term", "trade_scope", "group_scope",
           "links", "precedence", "parts", "gaps")
# Optional clause-level map of each file. Stored with the row, but kept out of --index so a
# long map only enters the session for the document being checked or the lines that match.
CONTENTS = "contents"
ROOT = WORK / "review-table"
ACTIVE = ROOT / "active.json"
SOURCE = WORK / "review-source.json"
VERSION = "analyse-review-index-v2"


def digest(path):
    value = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def find_review_table(pile, explicit=None):
    """Only the reserved Review_Table name is auto-discovered; ambiguity stops preparation."""
    if explicit:
        path = Path(explicit).expanduser().resolve()
        if not path.is_file() or path.is_symlink() or path.suffix.lower() not in {".csv", ".xlsx"}:
            raise ValueError("--review-table must name a regular CSV or XLSX file")
        return path
    candidates = [p for p in Path(pile).rglob("*") if p.is_file() and not p.is_symlink()
                  and p.stem.casefold() == "review_table" and p.suffix.lower() in {".csv", ".xlsx"}]
    if len(candidates) > 1:
        raise ValueError("Multiple Review_Table files found; pass --review-table with the intended file")
    return candidates[0].resolve() if candidates else None


def header(value):
    return re.sub(r"[\s-]+", "_", str(value or "").strip().casefold())


def read_table(path, sheet=None):
    """Read literal exported answers. Reject formulas and ambiguous sheets/headers."""
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            candidates = [s for s in book.worksheets if any(
                any(v is not None for v in row) for row in s.iter_rows(values_only=True))]
            if sheet:
                if sheet not in book.sheetnames:
                    raise ValueError(f"No worksheet named {sheet!r}")
                selected = book[sheet]
            elif len(candidates) == 1:
                selected = candidates[0]
            else:
                raise ValueError("Review_Table.xlsx needs one populated sheet, or an explicit --sheet")
            table = []
            for row in selected.iter_rows():
                if any(c.data_type == "f" for c in row):
                    raise ValueError("Review_Table contains formulas; export literal answer values first")
                table.append(["" if c.value is None else str(c.value) for c in row])
        finally:
            book.close()
    elif path.suffix.lower() == ".csv":
        raw = path.read_text(encoding="utf-8-sig")
        try:
            delimiter = csv.Sniffer().sniff(raw[:65536], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        table = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
    else:
        raise ValueError("Review_Table must be CSV (UTF-8) or XLSX")
    table = [row for row in table if any(str(v).strip() for v in row)]
    if not table:
        raise ValueError("Review_Table is empty")
    names = [header(v) for v in table[0]]
    while names and not names[-1]:
        names.pop()
    if not all(names) or len(set(names)) != len(names):
        raise ValueError("Review_Table has blank or duplicate column names")
    missing = set(("file_name",) + COLUMNS) - set(names)
    if missing:
        raise ValueError("Review_Table is missing columns: " + ", ".join(sorted(missing))
                         + ". Row 1 must use the short field names in inputs/Review_Table.example.csv; "
                         "full question headings are not mapped automatically. Rename the matching headers "
                         "in a copy of the export, preserving the original and cell values.")
    rows = []
    for number, values in enumerate(table[1:], 2):
        if len(values) > len(names) and any(str(v).strip() for v in values[len(names):]):
            raise ValueError(f"Review_Table row {number} has cells beyond the header")
        rows.append({name: str(values[i]) if i < len(values) else ""
                     for i, name in enumerate(names)})
    if not rows:
        raise ValueError("Review_Table has no document rows")
    return rows


def portable(value):
    return str(value).replace("\\", "/").casefold()


def match_rows(rows, inventory):
    """Resolve exact metadata, never fuzzy filenames or an LLM's guess."""
    by_name = {}
    for item in inventory:
        if item["doc_id"].isdigit():
            by_name.setdefault(portable(item["file_name"]), []).append(item)
    matched = {}
    for number, row in enumerate(rows, 2):
        candidates = by_name.get(portable(row["file_name"]), [])
        for key in ("doc_id", "sha256", "source_path"):
            if row.get(key):
                field = "original_path" if key == "source_path" else key
                wanted = row[key].zfill(3) if key == "doc_id" else row[key]
                candidates = [item for item in candidates if portable(item[field]) == portable(wanted)]
        if len(candidates) != 1:
            raise ValueError(f"Review_Table row {number}: filename/metadata match is "
                             f"{'ambiguous' if candidates else 'missing or contradictory'}; "
                             "use exact file_name plus doc_id, source_path or sha256")
        item = candidates[0]
        doc = item["doc_id"]
        if doc in matched:
            raise ValueError(f"Multiple Review_Table rows map to doc {doc}")
        if item["readable"] != "yes":
            raise ValueError(f"doc {doc} is not readable in the local inventory; resolve preparation first")
        for source in (Path(item["original_path"]), WORK_FILES / f'{doc}.{item["ext"]}'):
            if not source.is_file() or digest(source) != item["sha256"]:
                raise ValueError(f"doc {doc}: original/copy changed or missing; rerun /prepare and review")
        checked = COLUMNS + ((CONTENTS,) if CONTENTS in row else ())
        flags = [f"{column}: blank" for column in checked if not row[column].strip()]
        flags += [f"{column}: {token}" for column in checked
                  for token in sorted(set(re.findall(r"\b(?:NOT_REVIEWED|UNREADABLE|CONFLICTING)\b",
                                                    row[column])))]
        packet = {"schema": VERSION, "doc_id": doc, "inventory": item, "answers": row,
                  "flags": flags}
        packet["row_hash"] = hashlib.sha256(json.dumps(packet, sort_keys=True).encode()).hexdigest()
        matched[doc] = packet
    required = {r["doc_id"] for r in inventory if r["doc_id"].isdigit() and r["readable"] == "yes"}
    missing = required - matched.keys()
    if missing:
        raise ValueError("Review_Table has no row for readable documents: " + ", ".join(sorted(missing))
                         + ". Add rows (blank answers are allowed); no automatic source-reading fallback.")
    return matched


def archive_dependents(changed):
    if not changed:
        return
    archive = WORK / "history" / ("review-" + datetime.now().strftime("%Y%m%dT%H%M%S%f"))
    paths = [WORK / sub / f"{doc}{ext}" for doc in changed for sub, ext in
             (("cards", ".json"), ("cards", ".md"), ("forms", ".json"), ("filing", ".json"))]
    paths += [WORK / sub for sub in ("placements", "trees", "didnt-fit")]
    paths += [WORK / "logs/sort.csv", OUT]
    for path in paths:
        if path.exists():
            relative = Path("out") if path == OUT else path.relative_to(WORK)
            target = archive / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(path), str(target))


def import_table(path, sheet=None):
    inventory = read_csv(INVENTORY_CSV)
    if not inventory:
        raise ValueError("Run /prepare with the ERP and source corpus first")
    erp = read_json(WORK / "erp.json") or {}
    if not erp.get("confirmed"):
        raise ValueError("Confirm the ERP account column and side with /prepare before importing")
    if any(r["doc_id"].isdigit() and Path(r["original_path"]).resolve() == path.resolve()
           for r in inventory):
        raise ValueError("Review_Table was numbered as a contract; rerun /prepare with --review-table first")
    source_hash = digest(path)
    packets = match_rows(read_table(path, sheet), inventory)  # validate before any mutation
    old = read_json(ACTIVE) or {}
    changed = {d for d, packet in packets.items() if old.get("rows", {}).get(d) != packet["row_hash"]}
    changed |= set(old.get("rows", {})) - packets.keys()
    folder = ROOT / "imports" / source_hash
    folder.mkdir(parents=True, exist_ok=True)
    snapshot = folder / ("Review_Table" + path.suffix.lower())
    if path.resolve() != snapshot.resolve():
        shutil.copyfile(path, snapshot)
    if digest(snapshot) != source_hash or digest(path) != source_hash:
        raise ValueError("Review_Table changed during import; retry with a stable export")
    for doc, packet in packets.items():
        write_json(ROOT / "rows" / packet["row_hash"] / f"{doc}.json", packet)
    archive_dependents(changed)
    write_json(ACTIVE, {"schema": VERSION, "source_path": str(path.resolve()),
                       "source_hash": source_hash, "snapshot": str(snapshot), "sheet": sheet,
                       "imported_at": datetime.now(timezone.utc).isoformat(),
                       "rows": {d: p["row_hash"] for d, p in packets.items()}})
    write_json(SOURCE, {"source_path": str(path.resolve()), "sha256": source_hash, "sheet": sheet})
    say(f"Review_Table: {len(packets)} rows imported; {len(changed)} changed; no document readers run.")
    for doc, packet in packets.items():
        for flag in packet["flags"]:
            say(f"REVIEW doc {doc}: {flag}")


def packet_for(doc):
    active = read_json(ACTIVE) or {}
    row_hash = active.get("rows", {}).get(doc)
    packet = read_json(ROOT / "rows" / str(row_hash) / f"{doc}.json") if row_hash else None
    if packet is None:
        raise ValueError(f"No imported row for doc {doc}")
    return packet


def contents_lines(packet):
    """The stored map, one entry per line; None when the export has no contents column."""
    text = packet["answers"].get(CONTENTS)
    if text is None:
        return None
    return [line.strip() for line in text.splitlines() if line.strip()]


def index_view(packet, lookup=None):
    """Answers for --index/--lookup: every column literal except the map, which is summarised.

    With --lookup the matching map lines are returned so the session sees where a term occurs
    without loading the whole map.
    """
    answers = dict(packet["answers"])
    lines = contents_lines(packet)
    if lines is None:
        return answers
    doc = packet["doc_id"]
    answers[CONTENTS] = (f"{len(lines)} entries, {len(packet['answers'][CONTENTS])} characters; "
                         f"python scripts/review_table.py --contents {doc}")
    if lookup:
        answers[CONTENTS + "_matches"] = [line for line in lines if lookup.casefold() in line.casefold()]
    return answers


def show_contents(doc):
    packet = packet_for(doc)
    lines = contents_lines(packet)
    if lines is None:
        raise ValueError("This Review_Table has no contents column; ask the export for the clause map "
                         "or choose the location from the other cells")
    pages = packet["inventory"].get("pages", "")
    say(f"Contents map for doc {doc} (inventory pages: {pages or 'unknown'}; "
        f"{len(lines)} entries; a map is the export's claim, not a source read):")
    for line in lines:
        say(line)
    if not lines:
        say("(blank)")


def selected_ids(account=None):
    active = read_json(ACTIVE)
    if not active:
        raise ValueError("No imported Review_Table; run --import first")
    if not Path(active["source_path"]).is_file() or digest(active["source_path"]) != active["source_hash"]:
        raise ValueError("Review_Table changed or disappeared; reimport it before continuing")
    ids = set(active["rows"])
    if account:
        erp = read_json(WORK / "erp.json") or {}
        if account not in {a["account"] for a in erp.get("accounts", [])}:
            raise ValueError("Account is not an ERP row")
        ids = {r["doc_id"] for r in read_csv(WORK / "logs/sort.csv") if r["account"] == account}
        if not ids:
            raise ValueError("Account-only review needs an earlier matching run; run /analyse all first")
    inventory = {r["doc_id"]: r for r in read_csv(INVENTORY_CSV)}
    for doc in ids:
        packet = packet_for(doc)
        item = inventory.get(doc, {})
        if item.get("sha256") != packet["inventory"]["sha256"] or item.get("readable") != "yes":
            raise ValueError(f"doc {doc}: inventory is stale; reimport after /prepare")
        for path in (Path(item["original_path"]), WORK_FILES / f'{doc}.{item["ext"]}'):
            if not path.is_file() or digest(path) != item["sha256"]:
                raise ValueError(f"doc {doc}: source changed; rerun /prepare and review")
    return sorted(ids)


def display_records():
    """In-memory report fields, NOT sort cards. Never infer enums/dates from prose.

    Raw answers remain available in the index and report CSV. These conservative fields
    let the existing renderer display table-based judgments without model transcription.
    """
    records = {}
    for doc in selected_ids():
        answers = packet_for(doc)["answers"]
        record = {k: "not separately extracted (Review_Table)" for k in CARD_KEYS}
        record.update(doc_id=doc, q1_title=answers["document"] or f"doc {doc}",
                      q1_kind="other", q2_their_signing_entities="not found",
                      q2_their_group_companies="not found", q2_our_entity="not found",
                      q2_evidence=answers["parties"], q3_signed="can't tell",
                      q3_evidence=answers["execution"], q4_evidence=answers["term"],
                      q4_start_date="not found", q4_end="see Review_Table term",
                      q4_status="not assessed", q5_parts="can't tell",
                      q5_parts_detail="not found", q5_precedence=answers["precedence"],
                      q6_attaches_to="not found", q6_replaces="not found",
                      q6_referred_to_not_in_pile="not found", q6_evidence=answers["links"],
                      q7_trade_scope=answers["trade_scope"], q7_evidence=answers["trade_scope"],
                      q8_entities_covered=answers["group_scope"], q9_copy_or_draft_of="not found",
                      q10_oddities="Basis: Review_Table extraction. " + answers["gaps"])
        records[doc] = record
    return records


def assign(path, account=None):
    """Publish account decisions once, without rewriting any document's answers."""
    ids = set(selected_ids(account))
    inventory = {r["doc_id"]: r for r in read_csv(INVENTORY_CSV)}
    accounts = {r["account"] for r in read_json(WORK / "erp.json")["accounts"]}
    decisions = read_json(path)
    if not isinstance(decisions, list):
        raise ValueError("Assignments must be a JSON list of sort rows")
    produced, seen = [], set()
    for row in decisions:
        if not isinstance(row, dict) or any(not isinstance(v, str) for v in row.values()):
            raise ValueError("Each assignment must contain string fields")
        doc, target = row.get("doc_id", ""), row.get("account", "")
        if doc not in ids:
            raise ValueError(f"Assignment doc {doc} is outside the selected index")
        holding = target == "_no-name-found" or any(
            target.startswith(prefix) and bool(target[len(prefix):].strip())
            for prefix in ("_not-sure/", "_not-on-the-list/"))
        if target not in accounts and not holding:
            raise ValueError(f"Assignment account {target!r} is not an ERP row or holding target")
        if not row.get("basis", "").strip() or row.get("confidence") not in {"sure", "fairly sure", "not sure"}:
            raise ValueError(f"doc {doc}: give a match basis and sure/fairly sure/not sure confidence")
        if "known group" in row["basis"].lower() and row["confidence"] == "sure":
            raise ValueError(f"doc {doc}: a known group match is at most fairly sure; "
                             "record that human confirmation is needed")
        if (doc, target) in seen:
            raise ValueError(f"Duplicate assignment for doc {doc}, {target}")
        seen.add((doc, target))
        produced.append({**{k: row.get(k, "") for k in SORT_LOG_COLUMNS},
                         "original_path": inventory[doc]["original_path"]})
    if {r["doc_id"] for r in produced} != ids:
        raise ValueError("Assignments must cover every selected index row, including unresolved documents")
    previous = read_csv(WORK / "logs/sort.csv")
    kept = [r for r in previous if r["doc_id"] not in ids] if account else []
    if not account:
        kept = [{"doc_id": d, "original_path": r["original_path"], "account": "_unreadable",
                 "note": r.get("note", "")} for d, r in inventory.items()
                if d.isdigit() and r["readable"] != "yes"]
    # Reassignment invalidates old judgments for shared documents as well. Preserve them.
    if sorted(previous, key=lambda r: (r["doc_id"], r["account"])) != sorted(
            produced + kept, key=lambda r: (r["doc_id"], r["account"])):
        affected = {r["account"] for r in previous + produced if r["doc_id"] in ids} & accounts
        archive = WORK / "history" / ("review-match-" + datetime.now().strftime("%Y%m%dT%H%M%S%f"))
        for name in affected:
            for ext in (".csv", ".md"):
                source = WORK / "placements" / (safe_folder_name(name) + ext)
                if source.exists():
                    archive.mkdir(parents=True, exist_ok=True)
                    shutil.move(str(source), str(archive / source.name))
    write_csv(WORK / "logs/sort.csv", sorted(produced + kept, key=lambda r: (r["doc_id"], r["account"])),
              SORT_LOG_COLUMNS)
    say(f"Published {len(produced)} account assignments; no cards or source reads created.")


CHECKS = ROOT / "source-checks.jsonl"


def check_events():
    return [json.loads(line) for line in CHECKS.read_text(encoding="utf-8").splitlines()
            if line.strip()] if CHECKS.exists() else []


def current_checks(doc=None):
    active = read_json(ACTIVE) or {}
    return [r for r in check_events() if (not doc or r["doc_id"] == doc)
            and active.get("rows", {}).get(r["doc_id"]) == r["row_hash"]]


def log_check(path, finish=False):
    """Record access intent BEFORE source access, and findings afterwards. No implicit read."""
    data = read_json(path)
    if not isinstance(data, dict):
        raise ValueError("Source check must be a JSON object")
    if finish:
        request = next((r for r in current_checks() if r["check_id"] == data.get("check_id")
                        and r["event"] == "requested"), None)
        if request is None:
            raise ValueError("Source check has no current request")
        data = {**request, **data, "doc_id": request["doc_id"], "row_hash": request["row_hash"]}
        if any(r["check_id"] == data["check_id"] and r["event"] == "completed" for r in current_checks()):
            raise ValueError("This source check already has a recorded outcome")
        if not isinstance(data.get("finding"), str) or not data["finding"].strip():
            raise ValueError("Record the finding, including unresolved outcomes")
    for key in ("doc_id", "location", "reason", "mode"):
        if not isinstance(data.get(key), str) or not data[key].strip():
            raise ValueError(f"Source check requires {key}")
    if data["mode"] not in {"text", "image"}:
        raise ValueError("Source check mode must be text or image")
    if data["doc_id"] not in selected_ids():
        raise ValueError("Source check doc_id is not in the index")
    packet = packet_for(data["doc_id"])
    data.update(row_hash=packet["row_hash"], event="completed" if finish else "requested",
                recorded_at=datetime.now(timezone.utc).isoformat())
    if not finish:
        data["check_id"] = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%f")
    CHECKS.parent.mkdir(parents=True, exist_ok=True)
    with CHECKS.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(data, ensure_ascii=False) + "\n")
    say(json.dumps(data, ensure_ascii=False))
    if not finish:
        doc = data["doc_id"]
        say(f'Read only {data["location"]}: work/files/{doc}.{packet["inventory"]["ext"]} '
            f'or the matching excerpt from work/text/{doc}.txt. Then record the outcome.')
        if contents_lines(packet) is not None:
            show_contents(doc)  # the locator arrives when it is needed, not on every --index


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument("--import", dest="import_path", nargs="?", const="auto", help="CSV/XLSX, or prepared review source")
    action.add_argument("--index", action="store_true", help="print raw rows and any logged source checks")
    action.add_argument("--status", action="store_true", help="check identity/freshness and report source-check counts")
    action.add_argument("--assign", metavar="JSON", help="publish orchestrator account decisions")
    action.add_argument("--request-check", metavar="JSON", help="log a material doubt before targeted source access")
    action.add_argument("--finish-check", metavar="JSON", help="log actual scope and findings after source access")
    action.add_argument("--lookup", metavar="TEXT", help="search table answers and identities, without opening sources")
    action.add_argument("--contents", metavar="DOC", help="print one document's clause map from the table")
    parser.add_argument("--sheet", help="XLSX worksheet name")
    parser.add_argument("--account", help="scope to a previously matched ERP account")
    args = parser.parse_args()
    try:
        if args.import_path:
            path = args.import_path
            if path == "auto":
                configured = read_json(SOURCE) or read_json(ACTIVE) or {}
                path = configured.get("source_path")
                args.sheet = args.sheet or configured.get("sheet")
                if not path:
                    raise ValueError("No Review_Table registered by /prepare; pass --import PATH")
            import_table(Path(path).expanduser().resolve(), args.sheet)
        elif args.assign:
            assign(Path(args.assign), args.account)
        elif args.request_check or args.finish_check:
            log_check(Path(args.request_check or args.finish_check), finish=bool(args.finish_check))
        elif args.contents:
            doc = args.contents.zfill(3)
            if doc not in selected_ids(args.account):
                raise ValueError(f"doc {doc} is not in the index")
            show_contents(doc)
        else:
            ids = selected_ids(args.account)
            if args.index or args.lookup:
                for doc in ids:
                    packet = packet_for(doc)
                    if args.lookup and args.lookup.casefold() not in json.dumps(packet["answers"]).casefold():
                        continue
                    say(json.dumps({"doc_id": doc, "sha256": packet["inventory"]["sha256"],
                                    "answers": index_view(packet, args.lookup), "flags": packet["flags"],
                                    "source_checks": current_checks(doc)}, ensure_ascii=False))
            else:
                checks = [r for r in current_checks() if r["doc_id"] in ids]
                requested = [r for r in checks if r["event"] == "requested"]
                completed = [r for r in checks if r["event"] == "completed"]
                say(f"Review_Table: {len(ids)} index rows ready; no card conversion required.")
                say(f"Source checks: {len(requested)} requested, {len(completed)} completed, "
                    f"{len(requested) - len(completed)} pending; "
                    f"{len({r['doc_id'] for r in requested})} distinct documents requested.")
    except (ValueError, OSError, ImportError, csv.Error, zipfile.BadZipFile) as err:
        fail(str(err))
    return 0


if __name__ == "__main__":
    sys.exit(main())
